# -*- coding: utf-8 -*-
# extract keys and values from the it  and en languages files of a react application
# and build a xlsx file
import os, sys, re, logging, tempfile, datetime, json, glob
from openpyxl import Workbook, load_workbook
import subprocess
from io import BytesIO
from operator import itemgetter, attrgetter

Log = logging.getLogger(__name__)
Log.error(sys.argv)


# clever merge, coping with missing keys in either dictionary. only allowed languages: EN (column 3) and IT (column 4)
def merge_me(dict1,dict2):
    new_dict = {}
    both = dict1 | dict2
    # print(both)
    for k in both.keys():
        is_str=False
        if k in dict1:
            dict1_v = dict1[k]
            if isinstance(dict1_v, str):
                is_str=True
            else:
                valid_dict=dict1_v
        else:
            dict1_v = None
        if k in dict2:
            dict2_v = dict2[k]
            if isinstance(dict2_v, str):
                is_str=True
            else:
                valid_dict=dict2_v
        else:
            dict2_v = None

        if is_str:
            # solo Level1
            new_dict[k] = {"en": dict1_v, "it": dict2_v}
                # print ("L1: ", new_dict[k])
        else:
            nested = {"L2": True}
            for k2, v2 in valid_dict.items():
                # print("ITEMS: ",k2,v2)
                if k2 in dict1[k]:
                    en_inner_v = dict1[k][k2]
                else:
                    en_inner_v = None
                if k2 in dict2[k]:
                    it_inner_v = dict2[k][k2]
                else:
                    it_inner_v = None

                nested[k2] = { "en": en_inner_v, "it": it_inner_v }

            # print("NESTED: ", k, ": ", nested)
            new_dict[k] = nested

    # print("FINITO")
    return new_dict
    
def i18n2xls():
    dirname=sys.argv[1]
    tablename=sys.argv[2]
    try:

        report = ""
        trans={}
        out_content = BytesIO()
        wb = Workbook()
        ws = wb.active
        files = glob.glob(f"{dirname}/*/{tablename}.json")
        langs=[]
        for f in files:
            m=re.match(".*/(..)/.*.json",f)
            if m is not None:
                lang=m[1]
                print (lang)
                print (f)
                langs.append(lang)
                with open(f) as fp:
                    trans[lang]=json.loads(fp.read())
                    print("trans")

        # print(json.dumps(trans, indent=2))
        dict1=trans['en']
        dict2=trans['it']
        # print(dict1)
        merged = merge_me(dict1, dict2)

        ws.cell(row=1,column=1).value =  "L1"
        ws.cell(row=1,column=2).value =  "L2"
        ws.cell(row=1,column=3).value =  "EN"
        ws.cell(row=1,column=4).value =  "IT"

        row=1
        for k,v in merged.items():
            row+=1
            # print(f"KV {k=}, {v=}")
            ws.cell(row=row, column = 1).value = k
            if "L2" in v:
                v.pop('L2')
                print(f"L2 {k=} {v=}")
                for k2, v2 in v.items():
                    row+=1
                    ws.cell(row=row, column = 2).value = k2
                    ws.cell(row=row, column = 3).value = v2['en']
                    ws.cell(row=row, column = 4).value = v2['it']

            else:
                # print(f"{v=}")
                # remember: set elimina i duplicati
                ws.cell(row=row, column = 3).value = v['en']
                ws.cell(row=row, column = 4).value = v['it']

        wb.save(f"{tablename}.xlsx")
    except Exception as e:
        Log.error(f"(error: {e}")

if __name__ == "__main__":
    i18n2xls()
