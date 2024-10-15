# -*- coding: utf-8 -*-
# from a xlsx file extract the translations in locales/<language>/<tablename>.json
import os, sys, re, logging, tempfile, datetime, json
import csv
from openpyxl import Workbook, load_workbook
import subprocess
from operator import itemgetter, attrgetter

Log = logging.getLogger(__name__)
Log.error(sys.argv)

def xls2i18n():
    Log.error("xls2i18n")
    pulldown=False
    if sys.argv[1] == "--pulldown":
        pulldown=True
        fname=sys.argv[2]
    else:
        fname=sys.argv[1]

    dirname, tail = os.path.split(fname)
    basename = tail.split(".")[0]
    xlsx_name = f"{dirname}/{basename}.xlsx"
    Log.error(f"dirname: {dirname}, basename {basename}")
    try:

        if pulldown:
            wb = Workbook()
            ws = wb.active
            ws.title = "Sheet"
            data = open(fname)
            csv_data = list(csv.reader(data)) #Method used to open and read a csv file
            for i in csv_data:
                ws.append(i)
            wb.save(xlsx_name)

        report = ""
        xlsx = load_workbook(filename=xlsx_name)['Sheet']
       
        header = xlsx[1]
        langs = [ header[2].value, header[3].value]
        if len(header) ==5:
            langs.append(header[4].value)
        if len(header) ==6:
            langs.append(header[5].value)
        idx=1

        for lang in langs:
            lower_lang=lang.lower()
            new_L2_key=None
            idx+=1
            for in_rows in [xlsx]:

                label = in_rows[2][idx].value
                Log.error(f"TITLE: {label}")
                trans_dict = {}
                L2_dict={}
                for s in in_rows.iter_rows(min_row=3):
                    Log.error(s)
                    L1=s[0].value
                    L2=s[1].value
                    main_value=s[2].value
                    cur_value=s[idx].value
                    # Log.error(f"For s: {L1=} {L2=}, {main_value=} {cur_value=}")

                    if L1 is not None and L2 is not None and L2 == "L2":
                        # inizia un livello 2
                        print ("Inizio L2, ", L1)
                        if new_L2_key is not None:
                            # sputo fuori l'attuale L2  e reinizializzo
                            trans_dict[new_L2_key] = L2_dict
                        # Per i pull down bilivello posso nazionalizzare anche il nome del gruppo
                        L2_dict={}
                        new_L2_key = L1
                        L2_dict["Label"]=cur_value
                        
                    elif  L1 is None and L2 is None:
                        break

                    elif L2 is None: 
                        # livello 1
                        print(cur_value)
                        trans_dict[L1] = cur_value

                    elif  L1 is None:
                        # siamo dentro un L2 
                        L2_dict[L2] = cur_value

                    else:
                        Log.error(f"ELSE {L1}, {L2}, main value {main_value}, cur_value {cur_value}" )

                if new_L2_key is not None:
                    # ho ancora un dict in canna
                    print(f"esco {new_L2_key}")
                    trans_dict[new_L2_key] = L2_dict

            Log.error(trans_dict)
            if pulldown:
                out_dict = {f"{basename}_array": trans_dict}
                out_dict['Label'] = label
            else:
                out_dict = trans_dict

            # jtrans=json.dumps({f"{basename}_array": trans_dict},ensure_ascii=False,indent=2)
            jtrans=json.dumps(out_dict,ensure_ascii=False,indent=2)
            # no pun intended
            bjtrans=jtrans.encode('utf-8')
            out_dir=f"locales/{lower_lang}"
            try:
                os.makedirs(out_dir)
            except FileExistsError:
                Log.error(f"{out_dir} already exists")

            with open (f"{dirname}/{out_dir}/{basename}.json", "wb") as f:
                f.write(bjtrans)

    except Exception as e:
        Log.error(f"xls2i18n error: {e}")

if __name__ == "__main__":
    xls2i18n()

